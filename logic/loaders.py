from docx import Document as DocxDocument
from langchain_core.documents import Document
from pathlib import Path
from pypdf import PdfReader
import os
import pdf2image
import pytesseract

pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


def load_pdf_text(path: str) -> list[Document]:

    reader = PdfReader(path)

    pdf_documents: list[Document] = []

    for page_number, page in enumerate(reader.pages, start=1):

        text = page.extract_text().strip()

        if text:
            pdf_documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source_file":os.path.basename(path),
                        "page":page_number
                    }
                )
            )
        

    return pdf_documents

def load_pdf_scanned(path: str) -> list[Document]:
    images = pdf2image.convert_from_path(path)

    pdf_documents: list[Document] = []
    tessdata_dir = Path(__file__).resolve().parent.parent / "tessdata"
    os.environ["TESSDATA_PREFIX"] = str(tessdata_dir)

    for page_number, image in enumerate(images, start=1):
        text = pytesseract.image_to_string(
            image, 
            lang="deu", 
        ).strip()

        if text:
            pdf_documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source_file":os.path.basename(path),
                        "page":page_number
                    }
                )
            )

    return pdf_documents


def load_docx(path: str) -> list[Document]:
    file = DocxDocument(path)

    text = "\n".join(paragraph.text.strip() for paragraph in file.paragraphs).strip()

    if not text:
        return []

    return [
        Document(
            page_content=text,
            metadata={
                "source_file": os.path.basename(path),
                "page": 1
            }
        )
    ]


def load_file(path: str) -> list[Document]:
    file_format = Path(path).suffix.lower()

    if file_format == ".docx":
        return load_docx(path)
    elif file_format == ".pdf":
        pdf_docs = load_pdf_text(path)
        text = "\n".join(doc.page_content.strip() for doc in pdf_docs)

        if len(text) < 50:
            return load_pdf_scanned(path)
        else:
            return pdf_docs

    raise ValueError(f"Unsupported file type: {path}")
